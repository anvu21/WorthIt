"""Turn the hand-maintained Worth It workbook into clean restaurant records.

The workbook has two relevant sheets:
  Map  — one row per restaurant (the authoritative list of places)
  Main — one row per menu item, joined back to Map by restaurant name

Columns are located by header name so the sheet can be reordered freely.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

HYPERLINK_RE = re.compile(r'^=HYPERLINK\(\s*"([^"]*)"\s*,\s*"([^"]*)"\s*\)\s*$', re.IGNORECASE)
CLOSED_RE = re.compile(r"^\s*closed\b[\s:\-]*", re.IGNORECASE)
CLOSED_SUFFIX_RE = re.compile(r"\s*[(\[]\s*(permanently\s+)?closed\s*[)\]]\s*$", re.IGNORECASE)
REPLACEMENT_CHAR = "\ufffd"

TIER_MAP = {
    "$": "$",
    "$$": "$$",
    "$$$": "$$$",
    "side trip": "side",
    "one stop": "special",
    "holiday special": "special",
}
TIER_ORDER = ["$", "$$", "$$$", "side", "special"]

COUNTRY_ALIASES = {
    "united states": "USA",
    "united states of america": "USA",
    "us": "USA",
    "usa": "USA",
    "south korea": "Korea",
    "republic of korea": "Korea",
}

VOTER_COLUMNS = ["adam", "andrew", "steven", "others"]


@dataclass
class Appearance:
    title: str | None
    season: str | None
    episode: int | None
    tier: str | None
    video: str | None
    published: str | None

    def to_json(self) -> dict:
        return {
            "title": self.title,
            "season": self.season,
            "episode": self.episode,
            "tier": self.tier,
            "video": self.video,
            "published": self.published,
        }


@dataclass
class Item:
    name: str
    cost: float | None
    episode: str | None = None
    voters: list[str] = field(default_factory=list)

    def to_json(self) -> dict:
        return {"name": self.name, "cost": self.cost, "episode": self.episode, "voters": self.voters}


@dataclass
class Restaurant:
    row: int
    name: str
    closed: bool
    chain: bool | None
    address: str
    website: str | None
    notes: str | None
    appearances: list[Appearance]
    items: list[Item] = field(default_factory=list)
    city: str | None = None
    country: str | None = None
    lat: float | None = None
    lng: float | None = None
    coords_from_sheet: bool = False

    @property
    def tiers(self) -> list[str]:
        found = {a.tier for a in self.appearances if a.tier}
        return [t for t in TIER_ORDER if t in found]


@dataclass
class ParseResult:
    restaurants: list[Restaurant]
    warnings: list[str]


# --------------------------------------------------------------------------- helpers


def normalize_header(value) -> str:
    return re.sub(r"[^a-z0-9 ]", "", str(value or "").lower()).strip()


def norm_key(text: str | None) -> str:
    """Loose key used for matching names/addresses across sheets and imports."""
    text = (text or "").casefold()
    text = CLOSED_SUFFIX_RE.sub("", CLOSED_RE.sub("", text))
    text = re.sub(r",?\s*(united states|usa)\s*$", "", text)
    text = text.replace("&", " and ")
    return re.sub(r"[^0-9a-z\u00c0-\uffff]+", " ", text).strip()


def split_name(raw) -> tuple[str, str | None]:
    """Return (display name, hyperlink url) from a plain value or =HYPERLINK formula."""
    text = str(raw or "").strip()
    match = HYPERLINK_RE.match(text)
    if match:
        return match.group(2).strip(), match.group(1).strip() or None
    return text, None


def strip_closed(name: str) -> tuple[str, bool]:
    """'CLOSED Tokyo Dog', 'Bar Bolonat (CLOSED)', 'The Church Key [CLOSED]' → (name, True)."""
    if CLOSED_RE.match(name) or CLOSED_SUFFIX_RE.search(name):
        return CLOSED_SUFFIX_RE.sub("", CLOSED_RE.sub("", name)).strip(), True
    return name, False


def clean_url(url) -> str | None:
    text = str(url or "").strip()
    if not text or text.lower() in {"n/a", "-"}:
        return None
    if not re.match(r"^https?://", text, re.IGNORECASE):
        text = "https://" + text
    return text


def clean_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() in {"n/a", "-"}:
        return None
    return text


def fix_title(title: str | None) -> str | None:
    """The source file lost dashes/apostrophes to U+FFFD; repair the common shapes."""
    if not title:
        return title
    title = title.replace(f" {REPLACEMENT_CHAR} ", " – ")
    title = re.sub(f"{REPLACEMENT_CHAR}(?=s\\b)", "’", title)
    return title


def split_lines(value) -> list:
    """Split a multi-episode cell on newlines, dropping blank and '-' separator lines."""
    if value is None:
        return []
    if not isinstance(value, str):
        return [value]
    parts = [p.strip() for p in value.split("\n")]
    return [p for p in parts if p and p != "-"]


def pick(values: list, index: int):
    """Value for appearance `index`; a single value applies to every appearance."""
    if index < len(values):
        return values[index]
    if len(values) == 1:
        return values[0]
    return None


def parse_tier(value) -> str | None:
    return TIER_MAP.get(str(value or "").strip().lower())


def parse_season(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+(\.0+)?", text):
        return str(int(float(text)))
    return text or None


def parse_int(value) -> int | None:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def parse_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value or "").strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_cost(value) -> float | None:
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    match = re.search(r"\d[\d,]*(\.\d+)?", str(value or ""))
    return round(float(match.group(0).replace(",", "")), 2) if match else None


def parse_yes_no(value) -> bool | None:
    text = str(value or "").strip().lower()
    if text in {"yes", "y", "true"}:
        return True
    if text in {"no", "n", "false"}:
        return False
    return None


def parse_float(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_country(value) -> str | None:
    text = clean_text(value)
    if not text:
        return None
    return COUNTRY_ALIASES.get(text.lower(), text)


def display_address(address: str) -> str:
    return re.sub(r",\s*(United States|USA)\s*$", "", address.strip())


def parse_item_line(line: str) -> Item | None:
    """Parse a Map-sheet 'Items ordered' line like '$18.00 - Kinmedai White Ponzu'."""
    line = line.strip()
    if not line:
        return None
    match = re.match(r"^\$?\s*([\d,]+(?:\.\d+)?)\s*-?\s*(.+)$", line)
    if match:
        return Item(name=match.group(2).strip(" -"), cost=parse_cost(match.group(1)))
    return Item(name=line, cost=None)


# --------------------------------------------------------------------------- sheets


class Sheet:
    """Row access by header name, reading formulas and cached values side by side."""

    def __init__(self, formula_ws, value_ws):
        rows_f = list(formula_ws.iter_rows(values_only=True))
        rows_v = list(value_ws.iter_rows(values_only=True))
        self.headers = {normalize_header(h): i for i, h in enumerate(rows_f[0]) if h}
        self.rows = list(zip(rows_f[1:], rows_v[1:]))

    def col(self, *names: str) -> int | None:
        for name in names:
            if name in self.headers:
                return self.headers[name]
        return None

    def require(self, *names: str) -> int:
        index = self.col(*names)
        if index is None:
            raise ValueError(f"Missing required column: {names[0]!r}")
        return index


def get(row: tuple, index: int | None, cached: bool = False):
    """Read a cell; `row` is (formula_values, cached_values)."""
    if index is None:
        return None
    formulas, values = row
    value = formulas[index] if index < len(formulas) else None
    if cached or (isinstance(value, str) and value.startswith("=") and not HYPERLINK_RE.match(value)):
        return values[index] if index < len(values) else None
    return value


def parse_map_sheet(sheet: Sheet, warnings: list[str]) -> list[Restaurant]:
    c_name = sheet.require("restaurant")
    c_title = sheet.col("featured in episode")
    c_price = sheet.col("price point")
    c_season = sheet.col("season")
    c_episode = sheet.col("episode")
    c_items = sheet.col("items ordered")
    c_link = sheet.col("restaurant link")
    c_address = sheet.require("address")
    c_video = sheet.col("video link")
    c_date = sheet.col("publish date")
    c_chain = sheet.col("chain")
    c_notes = sheet.col("notes")
    c_lat = sheet.col("lat", "latitude")
    c_lng = sheet.col("lng", "lon", "long", "longitude")

    restaurants = []
    for offset, row in enumerate(sheet.rows):
        excel_row = offset + 2
        raw_name = get(row, c_name)
        if not clean_text(raw_name):
            continue
        name, link_from_formula = split_name(raw_name)
        name, closed = strip_closed(name)
        address = clean_text(get(row, c_address))
        if not address:
            warnings.append(f"Map row {excel_row} ({name}): no address — skipped")
            continue

        titles = [fix_title(str(t)) for t in split_lines(get(row, c_title))]
        prices = split_lines(get(row, c_price))
        seasons = split_lines(get(row, c_season))
        episodes = split_lines(get(row, c_episode))
        videos = split_lines(get(row, c_video))
        dates = split_lines(get(row, c_date))
        count = max(len(titles), len(videos), len(seasons), 1)

        appearances = [
            Appearance(
                title=pick(titles, i),
                season=parse_season(pick(seasons, i)),
                episode=parse_int(pick(episodes, i)),
                tier=parse_tier(pick(prices, i)),
                video=clean_url(pick(videos, i)),
                published=parse_date(pick(dates, i)),
            )
            for i in range(count)
        ]

        restaurant = Restaurant(
            row=excel_row,
            name=name,
            closed=closed,
            chain=parse_yes_no(get(row, c_chain)),
            address=display_address(address),
            website=clean_url(get(row, c_link) or link_from_formula),
            notes=clean_text(get(row, c_notes)),
            appearances=appearances,
        )

        lat, lng = parse_float(get(row, c_lat, cached=True)), parse_float(get(row, c_lng, cached=True))
        if lat is not None and lng is not None:
            restaurant.lat, restaurant.lng, restaurant.coords_from_sheet = lat, lng, True

        restaurant.items = [
            item for line in str(get(row, c_items) or "").split("\n") if (item := parse_item_line(line))
        ]

        for label, value in (("name", name), ("address", address)):
            if REPLACEMENT_CHAR in value:
                warnings.append(f"Map row {excel_row}: {label} has a broken character (�): {value!r}")

        restaurants.append(restaurant)
    return restaurants


def attach_main_sheet(sheet: Sheet, restaurants: list[Restaurant], warnings: list[str]) -> None:
    """Replace Map-sheet item text with Main-sheet items, and pick up city/country."""
    c_name = sheet.require("restaurant")
    c_title = sheet.col("featured in episode")
    c_cost = sheet.col("cost")
    c_item = sheet.col("item")
    c_address = sheet.col("address")
    c_city = sheet.col("city")
    c_country = sheet.col("country")
    voter_cols = {v: sheet.col(v) for v in VOTER_COLUMNS}

    by_name: dict[str, list[Restaurant]] = {}
    for r in restaurants:
        by_name.setdefault(norm_key(r.name), []).append(r)

    main_items: dict[int, list[Item]] = {}
    unmatched: set[str] = set()
    for offset, row in enumerate(sheet.rows):
        name, _ = split_name(get(row, c_name))
        name, _ = strip_closed(name)
        if not clean_text(name):
            continue
        candidates = by_name.get(norm_key(name), [])
        if len(candidates) > 1:
            address_key = norm_key(clean_text(get(row, c_address)))
            candidates = [r for r in candidates if norm_key(r.address) == address_key] or candidates[:1]
        if not candidates:
            unmatched.add(name)
            continue
        restaurant = candidates[0]

        if restaurant.city is None:
            restaurant.city = clean_text(get(row, c_city))
        if restaurant.country is None:
            restaurant.country = normalize_country(get(row, c_country))

        item_name = clean_text(get(row, c_item))
        if not item_name:
            continue
        voters = [v.capitalize() for v, col in voter_cols.items() if col is not None and get(row, col, cached=True)]
        title = split_lines(get(row, c_title))
        main_items.setdefault(id(restaurant), []).append(
            Item(
                name=" / ".join(split_lines(item_name)),
                cost=parse_cost(get(row, c_cost, cached=True)),
                episode=fix_title(str(title[0])) if title else None,
                voters=voters,
            )
        )

    for restaurant in restaurants:
        if id(restaurant) in main_items:
            restaurant.items = main_items[id(restaurant)]

    if unmatched:
        warnings.append(
            f"Main sheet: {len(unmatched)} restaurant name(s) not found on Map sheet (items ignored): "
            + ", ".join(sorted(unmatched)[:10])
            + (" …" if len(unmatched) > 10 else "")
        )


def parse_workbook(path: Path) -> ParseResult:
    warnings: list[str] = []
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    if "Map" not in formulas.sheetnames:
        raise ValueError("Workbook has no 'Map' sheet")

    restaurants = parse_map_sheet(Sheet(formulas["Map"], values["Map"]), warnings)
    if "Main" in formulas.sheetnames:
        attach_main_sheet(Sheet(formulas["Main"], values["Main"]), restaurants, warnings)
    else:
        warnings.append("No 'Main' sheet — items come from the Map sheet's 'Items ordered' column")
    return ParseResult(restaurants=restaurants, warnings=warnings)
